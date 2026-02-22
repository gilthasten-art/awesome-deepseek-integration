from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas


def write_docx(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    doc = Document()
    doc.add_heading(title, level=0)
    for header, body in sections:
        doc.add_heading(header, level=1)
        doc.add_paragraph(body)
    doc.save(path)


def write_pdf(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    c = canvas.Canvas(str(path), pagesize=LETTER)
    width, height = LETTER
    y = height - 50

    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, title)
    y -= 30

    for header, body in sections:
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, y, header)
        y -= 18
        c.setFont("Helvetica", 10)
        for line in _wrap_text(body, 95):
            if y < 50:
                c.showPage()
                y = height - 50
            c.drawString(50, y, line)
            y -= 14
        y -= 8

    c.save()


def _wrap_text(text: str, max_chars: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = []
    length = 0
    for word in words:
        if length + len(word) + (1 if current else 0) > max_chars:
            lines.append(" ".join(current))
            current = [word]
            length = len(word)
        else:
            current.append(word)
            length += len(word) + (1 if current[:-1] else 0)
    if current:
        lines.append(" ".join(current))
    return lines
