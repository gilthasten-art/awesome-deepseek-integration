from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas


class DocumentBuilder:
    def __init__(self, output_dir: str | Path) -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def write_pdf(self, title: str, body: str, file_name: str) -> Path:
        path = self.output_dir / file_name
        pdf = canvas.Canvas(str(path), pagesize=LETTER)
        width, height = LETTER

        y = height - 60
        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(60, y, title)
        y -= 28

        pdf.setFont("Helvetica", 11)
        for line in body.splitlines():
            if y < 50:
                pdf.showPage()
                y = height - 50
                pdf.setFont("Helvetica", 11)
            pdf.drawString(60, y, line[:120])
            y -= 16

        pdf.save()
        return path

    def write_docx(self, title: str, body: str, file_name: str) -> Path:
        path = self.output_dir / file_name
        doc = Document()
        doc.add_heading(title, level=1)
        for line in body.splitlines():
            doc.add_paragraph(line)
        doc.save(path)
        return path
