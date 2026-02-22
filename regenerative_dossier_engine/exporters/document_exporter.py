from __future__ import annotations

from pathlib import Path


class DocumentExporter:
    """Exports structured content to PDF and DOCX."""

    def export_docx(self, title: str, content: str, output_path: Path) -> Path:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError("python-docx is required for DOCX export") from exc

        doc = Document()
        doc.add_heading(title, level=1)
        for section in content.split("\n\n"):
            doc.add_paragraph(section.strip())
        doc.save(output_path)
        return output_path

    def export_pdf(self, title: str, content: str, output_path: Path) -> Path:
        try:
            from reportlab.lib.pagesizes import LETTER
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError("reportlab is required for PDF export") from exc

        doc = SimpleDocTemplate(str(output_path), pagesize=LETTER)
        styles = getSampleStyleSheet()
        story = [Paragraph(f"<b>{title}</b>", styles["Title"]), Spacer(1, 12)]

        for line in content.split("\n"):
            safe_line = line if line.strip() else "&nbsp;"
            story.append(Paragraph(safe_line, styles["BodyText"]))
            story.append(Spacer(1, 6))

        doc.build(story)
        return output_path
