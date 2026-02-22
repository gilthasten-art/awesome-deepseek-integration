from __future__ import annotations

from pathlib import Path

from regenerative_dossier_engine.models import BusinessDossier


class StructuredDocumentRenderer:
    """Renders dossiers as structured PDF and DOCX deliverables."""

    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def render(self, dossier: BusinessDossier, stem: str) -> tuple[Path, Path]:
        docx_path = self.output_dir / f"{stem}.docx"
        pdf_path = self.output_dir / f"{stem}.pdf"
        self._render_docx(dossier, docx_path)
        self._render_pdf(dossier, pdf_path)
        return docx_path, pdf_path

    def _render_docx(self, dossier: BusinessDossier, path: Path) -> None:
        from docx import Document

        doc = Document()
        doc.add_heading(dossier.title, level=1)
        doc.add_paragraph(f"Generated on: {dossier.generated_on.isoformat()}")
        for section, content in dossier.sections.items():
            doc.add_heading(section, level=2)
            doc.add_paragraph(content)
        if dossier.appendices:
            doc.add_heading("Appendices", level=2)
            for appendix in dossier.appendices:
                doc.add_paragraph(appendix, style="List Bullet")
        doc.save(path)

    def _render_pdf(self, dossier: BusinessDossier, path: Path) -> None:
        from reportlab.lib.pagesizes import LETTER
        from reportlab.pdfgen import canvas

        page = canvas.Canvas(str(path), pagesize=LETTER)
        width, height = LETTER
        y = height - 50

        def line(text: str, step: int = 16) -> None:
            nonlocal y
            page.drawString(50, y, text)
            y -= step

        line(dossier.title, step=22)
        line(f"Generated on: {dossier.generated_on.isoformat()}")
        line("-" * 60)
        for section, content in dossier.sections.items():
            line(section, step=20)
            for content_line in content.split("\n"):
                if y < 70:
                    page.showPage()
                    y = height - 50
                line(content_line)
            line("")
        if dossier.appendices:
            line("Appendices", step=20)
            for appendix in dossier.appendices:
                line(f"• {appendix}")
        page.save()
